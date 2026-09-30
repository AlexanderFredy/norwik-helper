"""Отчёт «где не добавлены фото»: группировка, текст и Excel.

Отвечает на вопрос менеджера («Покажи, где не добавлены фото») про **новые** товары — по
умолчанию за три месяца. Тотального обхода каталога здесь нет намеренно: это отдельная
работа и отдельный инструмент (решение админа 30.09.2026).

**ТЕКСТ ИЛИ ФАЙЛ — РЕШАЕТ ДЛИНА, А НЕ МОДЕЛЬ.** Два поста Telegram (8 192 знака) — граница,
за которой список перестаёт читаться в чате и начинает мешать: его листают, а не читают.
Дальше — Excel, где по нему можно работать: отсортировать, отметить сделанное, отдать
человеку, который грузит фото.

**СПИСОК НЕ ПРОХОДИТ ЧЕРЕЗ МОДЕЛЬ.** Инструмент отдаёт ей одну строку с числами, а сам
список уходит менеджеру напрямую (`ToolExecutor.pending_*`). Иначе за каждый вопрос
платили бы выходными токенами за пересказ сотни строк — и получали бы пересказ, в котором
ссылки однажды разойдутся с настоящими.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

#: Предел одного сообщения Telegram и сколько их допустимо на список.
POST = 4096
MAX_POSTS = 2


@dataclass(frozen=True)
class Row:
    """Строка отчёта: товар без фото."""
    tm: str
    collection: str
    name: str
    url: str
    created: str = ""

    @property
    def where(self) -> str:
        return f"{self.tm} / {self.collection}"


def _ru(iso: str) -> str:
    """ГГГГ-ММ-ДД → ДД.ММ.ГГГГ. Внутри даты ходят в ISO (они сравниваются строками), а
    читает отчёт человек, и ISO в русском тексте спотыкает."""
    parts = (iso or "").split("-")
    return ".".join(reversed(parts)) if len(parts) == 3 else iso


def render(rows: list[Row], *, since: str, checked: int, marks: int,
           no_card: int = 0, failed: int = 0, scope: str = "") -> str:
    """Текст для чата. Группировка по «марка / коллекция», внутри — наименование и ссылка."""
    where = f" по марке «{scope}»" if scope else ""
    if not rows:
        head = f"Все новые товары{where} с фото."
    else:
        head = f"Без фото на сайте: {len(rows)} поз."

    lines = [head, "",
             f"Проверено{where}: {checked} новых позиций (с {_ru(since)}), марок {marks}."]

    last = ""
    for row in sorted(rows, key=lambda r: (r.tm, r.collection, r.name)):
        if row.where != last:
            lines += ["", row.where]
            last = row.where
        lines.append(f"— {row.name} — {row.url}")

    # НЕ ВЫГРУЗИВШИЕСЯ НЕ ПОКАЗЫВАЕМ СПИСКОМ (решение админа: искать то, чего на сайте ещё
    # нет, незачем), но ЧИСЛОМ называем: иначе «без фото 3» при сорока не доехавших
    # читается как «почти всё в порядке».
    if no_card:
        lines += ["", f"Ещё {no_card} поз. на сайт пока не выгрузились — их не проверяли."]
    if failed:
        lines += ["", f"⚠️ {failed} поз. проверить не удалось: сайт не ответил."]
    return "\n".join(lines)


def fits_chat(text: str) -> bool:
    """Влезает ли в два поста. Граница — правило отчёта, а не особенность чата, поэтому
    считается здесь, а не в обработчике: он всего лишь режет то, что ему дали."""
    return len(text) <= MAX_POSTS * POST


def to_excel(rows: list[Row], path: Path) -> Path:
    """Тот же список файлом. Колонки отдельными — чтобы по ним можно было сортировать и
    фильтровать; ссылка живой гиперссылкой, иначе по ней не перейти из Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.title = "Без фото"
    headers = ("Марка", "Коллекция", "Наименование", "Ссылка", "Заведён")
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in sorted(rows, key=lambda r: (r.tm, r.collection, r.name)):
        sheet.append([row.tm, row.collection, row.name, row.url, row.created])
        link = sheet.cell(row=sheet.max_row, column=4)
        link.hyperlink = row.url
        link.style = "Hyperlink"

    for column, width in zip("ABCDE", (22, 26, 58, 38, 12)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path
