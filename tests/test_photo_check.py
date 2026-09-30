"""«Покажи, где не добавлены фото» (решение админа 30.09.2026).

Вопрос менеджера к агенту, а не инвентаризация каталога: смотрятся только товары,
заведённые за последние месяцы, и только выгружаемые на сайт. Сам список уходит мимо
модели — ей возвращаются числа.
"""
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from src.agent.tools import ToolExecutor, _months_ago
from src.onec.client import NomItem, Nomenclature, TradeMark
from src.website_tool import photo_report, photos


def item(ref="T1", site_id="1001", created="2026-09-01", name="Ламинат Classen Дуб Авола",
         parent="Adventure WR 1290x193x8", size="1290x193x8", collection="",
         not_exported=False):
    return NomItem(
        ref=ref, id=site_id, name=name, article="", unit="м2", size=size,
        product_type="Ламинат", collection=collection, parent=parent,
        collection_ref="F1", alt_units={}, purchase=None, retail=None, rrc=None,
        site_name=name, not_exported=not_exported, created=created)


class MonthsAgoTest(unittest.TestCase):
    """Считаем по календарю: «месяц = 30 дней» ошибается на трое суток, а за сутки
    заводят товары."""

    def test_plain(self):
        self.assertEqual(_months_ago(3, date(2026, 9, 30)), "2026-06-30")

    def test_crosses_the_year(self):
        self.assertEqual(_months_ago(3, date(2026, 2, 15)), "2025-11-15")

    def test_missing_day_moves_to_the_first(self):
        """31 мая минус 3 месяца — 31 февраля, которого нет: граница обязана существовать
        в календаре."""
        self.assertEqual(_months_ago(3, date(2026, 5, 31)), "2026-03-01")

    def test_december_rolls_over(self):
        self.assertEqual(_months_ago(1, date(2026, 1, 31)), "2025-12-31")


class FakeHTTP:
    """Отвечает кодами по пути запроса — так и выглядит сайт для проверки."""

    def __init__(self, codes):
        self.codes = codes
        self.asked: list[str] = []

    def head(self, path):
        self.asked.append(path)
        if path not in self.codes:
            raise AssertionError(f"неожиданный запрос: {path}")
        code = self.codes[path]
        if isinstance(code, Exception):
            raise code
        return type("R", (), {"status_code": code})()


def checker(codes):
    obj = photos.PhotoChecker.__new__(photos.PhotoChecker)
    obj._client = FakeHTTP(codes)
    obj._workers = 2
    return obj


class CheckerTest(unittest.TestCase):

    def test_folder_403_means_photo(self):
        """403 — папка есть, листинг закрыт. Сверено на 25 живых карточках."""
        c = checker({"/images/products/1001/": 403})
        self.assertEqual(c.status("1001"), photos.HAS)
        self.assertEqual(len(c._client.asked), 1, "страницу не трогаем вовсе")

    def test_no_folder_but_card_means_no_photo(self):
        c = checker({"/images/products/1001/": 404, "/item/1001": 200})
        self.assertEqual(c.status("1001"), photos.NONE)

    def test_no_folder_no_card(self):
        """Товар ещё не выгрузился на сайт — это не «нет фото»."""
        c = checker({"/images/products/1001/": 404, "/item/1001": 404})
        self.assertEqual(c.status("1001"), photos.NO_CARD)

    def test_empty_id_is_not_a_request(self):
        c = checker({})
        self.assertEqual(c.status(""), photos.NO_CARD)

    def test_site_failure_is_not_absence(self):
        """Придуманная работа хуже ненайденной: сбой сети не выдаём за отсутствие фото."""
        c = checker({"/images/products/1001/": RuntimeError("оборвалось")})
        self.assertEqual(c.status("1001"), photos.FAILED)

    def test_batch_asks_each_id_once(self):
        c = checker({"/images/products/1001/": 403, "/images/products/1002/": 403})
        out = c.statuses(["1001", "1002", "1001", " 1002 "])
        self.assertEqual(out, {"1001": photos.HAS, "1002": photos.HAS})
        self.assertEqual(len(c._client.asked), 2)


def row(tm="Classen", collection="Adventure WR", name="Дуб Авола", site_id="1001"):
    return photo_report.Row(tm=tm, collection=collection, name=name,
                            url=photos.item_url(site_id), created="2026-09-01")


class ReportTest(unittest.TestCase):

    def test_grouped_by_mark_and_collection(self):
        text = photo_report.render([row(name="Дуб Авола"), row(name="Дуб Брен")],
                                   since="2026-06-30", checked=9, marks=2)
        self.assertEqual(text.count("Classen / Adventure WR"), 1, "заголовок группы один")
        self.assertIn("— Дуб Авола — https://www.norwik.ru/item/1001", text)

    def test_link_is_the_site_card(self):
        self.assertEqual(photos.item_url("189490"), "https://www.norwik.ru/item/189490")

    def test_nothing_missing_is_said_plainly(self):
        text = photo_report.render([], since="2026-06-30", checked=12, marks=3)
        self.assertIn("Все новые товары с фото", text)
        self.assertIn("12 новых позиций", text)

    def test_not_exported_are_counted_but_not_listed(self):
        """Показывать их админ просил не списком — но молчать о них нельзя: «без фото 3»
        при сорока не доехавших читается как «почти всё в порядке»."""
        text = photo_report.render([row()], since="2026-06-30", checked=44, marks=1,
                                   no_card=40)
        self.assertIn("40 поз. на сайт пока не выгрузились", text)
        self.assertEqual(text.count("https://"), 1, "их ссылок в списке нет")

    def test_site_failures_are_named(self):
        text = photo_report.render([], since="2026-06-30", checked=5, marks=1, failed=2)
        self.assertIn("2 поз. проверить не удалось", text)

    def test_two_posts_is_the_border(self):
        self.assertTrue(photo_report.fits_chat("x" * 8192))
        self.assertFalse(photo_report.fits_chat("x" * 8193))

    def test_excel_has_headers_and_rows(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            path = photo_report.to_excel([row(name="Дуб Авола")],
                                         Path(tmp) / "нет фото.xlsx")
            sheet = load_workbook(path).active
            self.assertEqual([c.value for c in sheet[1]][:4],
                             ["Марка", "Коллекция", "Наименование", "Ссылка"])
            self.assertEqual(sheet.cell(row=2, column=3).value, "Дуб Авола")
            self.assertEqual(sheet.cell(row=2, column=4).hyperlink.target,
                             "https://www.norwik.ru/item/1001")


class FakeOnec:
    def __init__(self, dumps, marks=None, errors=0):
        self._dumps, self._errors = dumps, errors
        self.marks = marks or [TradeMark(name="Classen / Классен", code="1")]
        self.asked: list[tuple] = []

    def selling_tm(self, all_marks=False):
        return list(self.marks)

    def by_tm_all(self, code, created_from=None, **kw):
        self.asked.append((code, created_from))
        items = self._dumps.get(code, [])
        return Nomenclature(tm="ТМ", total=len(items), items=list(items),
                            errors=[{"ref": "X"}] * self._errors)


class FakeChecker:
    def __init__(self, verdicts):
        self.verdicts = verdicts

    def __call__(self, *_a, **_kw):
        return self

    def statuses(self, ids):
        return {(i or "").strip(): self.verdicts.get(i, photos.HAS) for i in ids}

    def close(self):
        pass


class ToolTest(unittest.IsolatedAsyncioTestCase):

    def executor(self, onec):
        return ToolExecutor(mail=None, norwik=None, onec=onec)

    async def run_tool(self, onec, verdicts, inp=None):
        ex = self.executor(onec)
        with patch.object(photos, "PhotoChecker", FakeChecker(verdicts)):
            answer = await ex._photos_missing(inp or {})
        return ex, answer

    async def test_list_goes_to_the_manager_not_through_the_model(self):
        """Пересказ сотни строк — это выходные токены за копирование и риск, что ссылки
        разойдутся с настоящими."""
        onec = FakeOnec({"1": [item(site_id="1001"), item(site_id="1002")]})
        ex, answer = await self.run_tool(onec, {"1001": photos.NONE})
        self.assertIn("Найдено 1", answer)
        self.assertNotIn("norwik.ru", answer, "ссылок модель не видит")
        text, file = ex.take_pending()
        self.assertIn("https://www.norwik.ru/item/1001", text)
        self.assertIsNone(file)

    async def test_pending_is_cleared_once_taken(self):
        onec = FakeOnec({"1": [item()]})
        ex, _ = await self.run_tool(onec, {"1001": photos.NONE})
        ex.take_pending()
        self.assertEqual(ex.take_pending(), (None, None))

    async def test_date_filter_goes_to_1c(self):
        """Отбор на стороне 1С: иначе тянули бы каталог целиком, чтобы отбросить почти всё."""
        onec = FakeOnec({"1": [item()]})
        await self.run_tool(onec, {}, {"months": 2})
        self.assertEqual(onec.asked[0][1], _months_ago(2))

    async def test_old_items_are_dropped_even_if_1c_sent_them(self):
        """Старый by-tm.bsl параметр не заметит и отдаст всё — отбор повторяем у себя."""
        onec = FakeOnec({"1": [item(site_id="1001", created="2020-01-01"),
                               item(site_id="1002", created=_months_ago(1))]})
        ex, answer = await self.run_tool(onec, {"1001": photos.NONE, "1002": photos.NONE})
        text, _ = ex.take_pending()
        self.assertIn("item/1002", text)
        self.assertNotIn("item/1001", text)

    async def test_no_dates_at_all_is_an_error_not_an_empty_answer(self):
        """«Новых нет» читалось бы как «всё в порядке», а поле просто не выложено."""
        onec = FakeOnec({"1": [item(created=""), item(created="")]})
        _, answer = await self.run_tool(onec, {})
        self.assertIn("не отдаёт дату создания", answer)
        self.assertIn("by-tm.bsl", answer)

    async def test_discontinued_are_skipped(self):
        onec = FakeOnec({"1": [item(site_id="1001", not_exported=True)]})
        _, answer = await self.run_tool(onec, {"1001": photos.NONE})
        self.assertIn("Новых товаров", answer)

    async def test_items_without_a_card_are_not_listed(self):
        onec = FakeOnec({"1": [item(site_id="1001"), item(site_id="1002")]})
        ex, answer = await self.run_tool(
            onec, {"1001": photos.NO_CARD, "1002": photos.NONE})
        text, _ = ex.take_pending()
        self.assertNotIn("item/1001", text)
        self.assertIn("1 поз. на сайт пока не выгрузились", text)

    async def test_long_list_becomes_excel(self):
        rows = [item(ref=f"R{i}", site_id=str(2000 + i),
                     name=f"Ламинат Classen Очень Длинное Название Декора Номер {i}")
                for i in range(200)]
        onec = FakeOnec({"1": rows})
        ex, answer = await self.run_tool(
            onec, {str(2000 + i): photos.NONE for i in range(200)})
        text, file = ex.take_pending()
        self.assertIsNone(text, "в чат не шлём — не влезло")
        self.assertTrue(str(file).endswith(".xlsx"))
        self.assertIn("файлом Excel", answer)
        Path(str(file)).unlink(missing_ok=True)

    async def test_mark_filter(self):
        onec = FakeOnec({"1": [item()]},
                        marks=[TradeMark(name="Classen / Классен", code="1"),
                               TradeMark(name="Peli", code="2")])
        await self.run_tool(onec, {"1001": photos.HAS}, {"tm": "classen"})
        self.assertEqual([code for code, _ in onec.asked], ["1"])

    async def test_unknown_mark(self):
        onec = FakeOnec({})
        _, answer = await self.run_tool(onec, {}, {"tm": "нетакой"})
        self.assertIn("нет среди выгружаемых", answer)

    async def test_lost_positions_are_named_to_the_model(self):
        onec = FakeOnec({"1": [item()]}, errors=3)
        _, answer = await self.run_tool(onec, {"1001": photos.NONE})
        self.assertIn("не отдала 3", answer)

    async def test_without_1c(self):
        ex = self.executor(None)
        self.assertIn("не настроена", await ex._photos_missing({}))


if __name__ == "__main__":
    unittest.main()
